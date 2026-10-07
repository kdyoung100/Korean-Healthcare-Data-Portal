#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
논문DB.xlsx — 데이터셋별 PubMed 논문 DB (수집 → 검수 → 사이트 표시)

    python paper_db.py fetch                 # PubMed 에서 모아 논문DB.xlsx 에 합친다 (새 논문 = '미검수')
    python paper_db.py pending [파일.json]    # 미검수 논문을 JSON 으로 뽑는다 (검수 에이전트용, 초록 포함)
    python paper_db.py review 결정.json       # 검수 결과를 논문DB.xlsx 에 반영한다
    python paper_db.py stats                 # 상태별 건수
    python paper_db.py nosummary [파일.json]  # '확인' 인데 한줄요약이 빈 논문을 뽑는다 (초록 포함)
    python paper_db.py summary 요약.json      # 한줄요약을 반영한다  [{"pmid": "...", "summary": "..."}]
    python paper_db.py check                 # 동명이인 확인 대상을 '연구자확인' 시트에 채운다 (판정 칸은 보존)
    python paper_db.py waiting               # 담당자 승인을 기다리는 논문 (데이터셋별 건수 + 목록)
    python paper_db.py approve CD003 12345678 # 담당자 OK — 데이터셋ID(그 데이터의 '확인' 논문 전부) 또는 PMID
    python paper_db.py reject 12345678       # 담당자 반려 — 사이트에 안 나온다

  사이트에 나오는 조건 = 그 데이터셋에 대해 데이터셋판정 '확인'(AI·사람 검수) **그리고** 승인데이터셋에 그 데이터셋(담당자 OK).
  판정·승인은 논문 × 데이터셋 단위다 — 이미 승인된 논문에 fetch 가 새 데이터셋ID를 붙이면 그 데이터셋은
  다시 '미검수'(pending 에 나옴)이고, 검수·승인 전에는 그 데이터셋 화면에 나오지 않는다.

  검색식 : pubmed_queries.json  (데이터셋ID → PubMed 검색식)
  DB     : 논문DB.xlsx          (한 행 = 논문 1편. 커밋할 것 — CI 는 PubMed 를 부르지 않고 이 파일만 읽는다)
  사이트 : build.py 가 load_for_build() 로 읽어 data.json 의 papers 에 넣는다.

검수 결정 JSON 형식:
    [{"pmid": "12345678", "status": "확인", "reason": "초록에 '결핵 신고자료 이용' 명시"}, ...]
    status 는 확인 / 제외 / 보류 중 하나. reviewer 를 안 적으면 'AI 검수'.
    "summary" 를 함께 적으면 한줄요약 칸도 채운다(사이트의 논문 카드 설명, 한국어 1~2문장).

fetch 는 사람이 적은 검수 칸(검수상태·검수근거·한줄요약·검수자·검수일·메모)을 절대 덮어쓰지 않는다.
검색식에서 빠진 논문도 지우지 않는다(검수 기록 보존).
"""
import json, os, re, sys, time, urllib.parse, urllib.request
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

BASE = Path(__file__).resolve().parent
DB = BASE / "논문DB.xlsx"
QUERIES = BASE / "pubmed_queries.json"
EUTILS = "https://eutils.ncbi.nlm.nih.gov/entrez/eutils/"
PUBMED = "https://pubmed.ncbi.nlm.nih.gov/"

STATUSES = ["미검수", "확인", "제외", "보류"]
APPROVALS = ["승인", "반려"]       # 담당자승인 칸 — 비어 있으면 '승인 대기'. 사람(담당자)만 정한다
REVIEW_COLS = ["검수상태", "데이터셋판정", "담당자승인", "승인데이터셋", "승인일", "검수근거", "한줄요약", "검수자", "검수일", "메모"]     # fetch 가 건드리지 않는 칸
COLS = ["PMID", "데이터셋ID", "검수상태", "데이터셋판정", "담당자승인", "승인데이터셋", "승인일", "검수근거", "한줄요약", "검수자", "검수일",
        "제목", "저자", "저자수", "저널", "발행연도", "발행일", "DOI",
        "초록", "수집일", "메모"]
WIDTH = {"PMID": 11, "데이터셋ID": 14, "검수상태": 9, "데이터셋판정": 22, "담당자승인": 10, "승인데이터셋": 16, "승인일": 11, "검수근거": 40, "한줄요약": 50, "검수자": 11, "검수일": 11,
         "제목": 60, "저자": 30, "저자수": 7, "저널": 22, "발행연도": 8, "발행일": 12, "DOI": 24,
         "초록": 60, "수집일": 11, "메모": 24}
Q_COLS = ["데이터셋ID", "검색식", "PubMed 전체건수", "조회일"]
A_COLS = ["PMID", "순서", "이름", "약칭", "ORCID", "소속", "기관"]
FIX_COLS = ["PMID", "이름", "연구자ID", "메모"]      # '연구자보정' 시트 — 사람이 적는다. fetch 가 보존한다
CHK_COLS = ["확인ID", "연구자키", "이름", "사유", "논문수", "기관", "ORCID", "PMID", "대상 연구자키", "판정", "판정자", "판정일", "메모"]
CHK_KEEP = ["판정", "판정자", "판정일", "메모"]      # '연구자확인' 시트에서 사람이 적는 칸 — check 가 덮어쓰지 않는다
JUDGE = ["같은 사람", "다른 사람"]
FILL = {"확인": "E3F4E5", "제외": "F4E3E3", "보류": "FFF4D6"}
APPROVE_FILL = {"승인": "C8E6C9", "반려": "F4E3E3", "": "FFE0B2"}    # 빈칸(확인인데 승인 대기)은 주황


# ── PubMed 호출 ────────────────────────────────────────────────────
def _get(endpoint, **params):
    params.update(db="pubmed", tool="kmdp-portal")
    if os.environ.get("NCBI_API_KEY"):
        params["api_key"] = os.environ["NCBI_API_KEY"]
    if os.environ.get("NCBI_EMAIL"):
        params["email"] = os.environ["NCBI_EMAIL"]
    data = urllib.parse.urlencode(params).encode()      # POST — PMID 가 많아도 URL 길이 제한에 안 걸림
    for attempt in range(3):
        try:
            with urllib.request.urlopen(EUTILS + endpoint, data=data, timeout=60) as r:
                return r.read()
        except Exception:
            if attempt == 2:
                raise
            time.sleep(2 * (attempt + 1))
        finally:
            time.sleep(0.12 if "api_key" in params else 0.4)   # 초당 3건(키 있으면 10건) 한도


def _search(q):
    es = json.loads(_get("esearch.fcgi", term=q, retmax=9999, retmode="json"))["esearchresult"]
    return int(es.get("count", 0)), es.get("idlist", [])


# 소속 문자열에서 '기관'만 뽑는다. "Department of X, Seoul National University College of Medicine,
# Seoul, Korea" → "서울대학교". 자주 나오는 국내 기관은 한글명으로 바꾼다(INST_KO).
_INST_KEY = r"University|Hospital|Agency|Institute|Center|Centre|College|Ministry|Service|Corporation|Association|Foundation|Council|Organi[sz]ation|Office|Society|Command|Clinic|대학|병원"
_INST_SKIP = r"^(Department|Dept|Division|Section|Laboratory|Lab|Unit|Team|Program|Graduate School|School of|College of|Center for|Centre for|Institute of|Office of|Bureau)\b"
INST_KO = {
    "Korea Disease Control and Prevention Agency": "질병관리청",
    "Korea Centers for Disease Control and Prevention": "질병관리본부",
    "Korean Institute of Tuberculosis": "대한결핵협회 결핵연구원",
    "Korean National Tuberculosis Association": "대한결핵협회",
    "National Health Insurance Service": "국민건강보험공단",
    "Health Insurance Review and Assessment Service": "건강보험심사평가원",
    "Ministry of Health and Welfare": "보건복지부",
    "Seoul National University": "서울대학교", "Yonsei University": "연세대학교", "Korea University": "고려대학교",
    "Sungkyunkwan University": "성균관대학교", "Catholic University of Korea": "가톨릭대학교",
    "University of Ulsan": "울산대학교", "Ulsan University": "울산대학교", "Hallym University": "한림대학교",
    "Ewha Womans University": "이화여자대학교", "Chung-Ang University": "중앙대학교", "Hanyang University": "한양대학교",
    "Kyung Hee University": "경희대학교", "Inha University": "인하대학교", "Ajou University": "아주대학교",
    "Pusan National University": "부산대학교", "Kyungpook National University": "경북대학교",
    "Chonnam National University": "전남대학교", "Chungbuk National University": "충북대학교",
    "Chungnam National University": "충남대학교", "Konkuk University": "건국대학교",
    "Soonchunhyang University": "순천향대학교", "Inje University": "인제대학교", "Gachon University": "가천대학교",
    "National Cancer Center": "국립암센터", "National Medical Center": "국립중앙의료원",
    "Masan National Tuberculosis Hospital": "국립마산병원", "Asan Medical Center": "서울아산병원",
    "Samsung Medical Center": "삼성서울병원", "Severance Hospital": "세브란스병원",
    "Seoul National University Hospital": "서울대학교병원",
    "Seoul National University Bundang Hospital": "분당서울대학교병원",
    "International Tuberculosis Research Center": "국제결핵연구소",
    "Korea Centers for Disease Control & Prevention": "질병관리본부",
    "Incheon St Mary's Hospital": "인천성모병원", "Seoul St Mary's Hospital": "서울성모병원",
    "Daejeon St Mary's Hospital": "대전성모병원", "Uijeongbu St Mary's Hospital": "의정부성모병원",
    "Ilsan Paik Hospital": "일산백병원", "Inje University Ilsan Paik Hospital": "일산백병원",
    "Kangdong Sacred Heart Hospital": "강동성심병원", "Hallym University Kangdong Sacred Heart Hospital": "강동성심병원",
    "Korea University Guro Hospital": "고려대학교 구로병원", "Pusan National University Hospital": "부산대학교병원",
    "Pusan National University Yangsan Hospital": "양산부산대학교병원", "Ulsan University Hospital": "울산대학교병원",
    "Yongin Severance Hospital": "용인세브란스병원", "Konyang University": "건양대학교", "Dankook University": "단국대학교",
    "National Evidence-based Collaborating Agency": "한국보건의료연구원",
    "National Evidence-Based Healthcare Collaborating Agency": "한국보건의료연구원",
    "Korea University College of Health Science": "고려대학교",
    "Chonnam National University": "전남대학교", "Chonnam National University Medical School": "전남대학교",
    "Chonnam National University Hospital": "전남대학교병원", "Chonnam National University Hwasun Hospital": "화순전남대학교병원",
    "Yeungnam University": "영남대학교", "Yeungnam University Medical Center": "영남대학교병원",
    "Yeungnam University College of Medicine": "영남대학교", "Mokdong Hospital": "이대목동병원",
    "Ewha Womans University Mokdong Hospital": "이대목동병원", "Ewha Womans University Medical Center": "이화여자대학교의료원",
    "Hallym University Dongtan Sacred Heart Hospital": "동탄성심병원",
}
# 동일인 판정용 '기관 계열' — 병원은 소속 대학으로, 개편된 기관은 지금 이름으로 본다 (화면 표시는 그대로)
INST_FAMILY = {
    "질병관리본부": "질병관리청", "대한결핵협회 결핵연구원": "대한결핵협회", "이대목동병원": "이화여자대학교",
    "이화여자대학교의료원": "이화여자대학교", "세브란스병원": "연세대학교", "용인세브란스병원": "연세대학교",
    "서울대학교병원": "서울대학교", "분당서울대학교병원": "서울대학교", "서울성모병원": "가톨릭대학교",
    "인천성모병원": "가톨릭대학교", "대전성모병원": "가톨릭대학교", "의정부성모병원": "가톨릭대학교",
    "삼성서울병원": "성균관대학교", "서울아산병원": "울산대학교", "울산대학교병원": "울산대학교",
    "부산대학교병원": "부산대학교", "양산부산대학교병원": "부산대학교", "고려대학교 구로병원": "고려대학교",
    "일산백병원": "인제대학교", "강동성심병원": "한림대학교", "동탄성심병원": "한림대학교",
    "화순전남대학교병원": "전남대학교", "전남대학교병원": "전남대학교", "영남대학교병원": "영남대학교",
}


# 같은 기관의 다른 표기·부속병원 → 계열 (동명이인 판정이 표기 차이로 '기관 2곳'이 되지 않게)
FAMILY_ALIAS = {
    "Bongseng Memorial Hospital": "Bong Seng Memorial Hospital", "Dongkuk University": "Dongguk University",
    "SMG-SNU Boramae Medical Center": "서울대학교", "SMG-SNU Boramae Medical Centre": "서울대학교",
    "Ewha Woman's University": "이화여자대학교", "Ewha Women's University": "이화여자대학교",
    "Kangnam Sacred Heart Hospital": "한림대학교", "Gangnam Severance Hospital": "연세대학교",
    "Severance Children's Hospital": "연세대학교", "Severance Cardiovascular Hospital": "연세대학교",
    "Renji Hospital": "Shanghai Jiao Tong University", "Busan Paik Hospital": "인제대학교",
    "National Emergency Medical Center of National Medical Center": "국립중앙의료원",
    "Kyung Hee University Hospital at Gangdong": "경희대학교",
    "Catholic University": "가톨릭대학교", "Soon Chun Hyang University": "순천향대학교",
    "National Health Insurance Services": "국민건강보험공단", "National Health Insurance Service Ilsan Hospital": "국민건강보험공단",
    "Kangbuk Samsung Hospital": "성균관대학교", "Samsung Kangbuk Hospital": "성균관대학교",
    "Uijeongbu Eulji Medical Center": "Eulji University", "Unity Health Toronto": "St Michael's Hospital",
}
# 기관을 못 뽑고 부서·연구실 이름만 남은 것 — 동명이인 판정에서 '모름'으로 본다
_FAMILY_UNKNOWN = r"^(Department|Division|Section|Laboratory) of [^,]*$|^(University|Artificial Intelligence and Big Data Research Center|Mental Health Research Institute|Public Health and Medical Service Office)$"
_UNIV_KO = sorted(((k, v) for k, v in INST_KO.items() if k.endswith("University")), key=lambda t: -len(t[0]))


def inst_family(name):
    n = re.sub(r"\s*\([^)]*\)$", "", (name or "").strip())        # "… (KIDS)" 같은 끝 약칭은 떼고 본다
    n = FAMILY_ALIAS.get(n, n)
    if n in INST_FAMILY:
        return INST_FAMILY[n]
    m = re.match(r"^(\S+대학교)(병원|의료원|\s.*)?$", n)
    if m:
        return m.group(1)
    for k, v in _UNIV_KO:                  # "Department of Surgery Yonsei University" → 연세대학교
        if k in n:
            return INST_FAMILY.get(v, v)
    if re.match(_FAMILY_UNKNOWN, n, re.I):
        return ""
    m = re.match(r"^(University of [A-Z][\w-]*(?: [A-Z][\w-]*)*|.*?University)\b", n)
    f = m.group(1) if m else n
    return FAMILY_ALIAS.get(f, f)
# 한글 소속("질병관리청 감염병정책국 결핵정책과")은 첫 기관 단위까지만
_KO_ORG = r"^(\S*?(청|본부|대학교|대학|병원|공단|연구원|협회|의료원|센터|연구소))(\s|$)"


def institution(aff):
    aff = (aff or "").split(" / ")[0].replace("St.", "St")
    m = re.match(_KO_ORG, aff.strip())
    if m and re.search(r"[가-힣]", m.group(1)):
        return INST_KO.get(m.group(1), m.group(1))
    aff = re.sub(r"(Electronic address:|E-?mail:).*$", "", aff, flags=re.I)
    aff = re.sub(r"\S+@\S+", "", aff)
    parts = [x.strip(" .;") for x in re.split(r"[,;]", aff) if x.strip(" .;")]
    known = next((x for x in parts if x in INST_KO or re.sub(r"^The\s+", "", x) in INST_KO), None)
    if known:
        return INST_KO[re.sub(r"^The\s+", "", known)]
    if any(x == "Centers for Disease Control and Prevention" for x in parts) and re.search(r"Korea|Cheongju|Osong", aff):
        return "질병관리본부"
    cands = [x for x in parts if re.search(_INST_KEY, x, re.I) and not re.search(_INST_SKIP, x, re.I)]
    pick = next((x for x in cands if re.search(r"University|Hospital|Agency|Ministry|Service|대학|병원", x, re.I)), None) \
        or (cands[0] if cands else (parts[0] if parts else ""))
    pick = re.sub(r"\s+(College|School|Graduate School) of (Medicine|Public Health|Nursing|Pharmacy)\b.*$", "", pick, flags=re.I)
    pick = re.sub(r"^The\s+", "", pick)
    head = re.sub(r"\s+and\s+.*$", "", pick)            # 'Yeungnam University and Regional Center …' → 앞 기관
    pick = head if head in INST_KO else pick
    return INST_KO.get(pick, pick)


def _orcid(s):
    """'https://orcid.org/0000-0002-1825-0097' · '0000000218250097' → '0000-0002-1825-0097'"""
    d = re.sub(r"[^0-9Xx]", "", (s or "").split("orcid.org/")[-1]).upper()
    return "-".join(d[i:i + 4] for i in range(0, 16, 4)) if len(d) == 16 else ""


def _details(pmids):
    """PMID → 서지정보 + 초록. efetch XML 한 번으로 다 나온다."""
    out = {}
    for i in range(0, len(pmids), 200):
        root = ET.fromstring(_get("efetch.fcgi", id=",".join(pmids[i:i + 200]), retmode="xml"))
        for art in root.iter("PubmedArticle"):
            txt = lambda el: "".join(el.itertext()).strip() if el is not None else ""
            pmid = txt(art.find(".//MedlineCitation/PMID"))
            a = art.find(".//Article")
            authors, detail = [], []
            for au in a.findall(".//AuthorList/Author"):
                last, ini, fore = au.findtext("LastName"), au.findtext("Initials"), au.findtext("ForeName")
                name = f"{last} {ini}".strip() if last else au.findtext("CollectiveName")
                if not name:
                    continue
                authors.append(name)
                aff = " / ".join(txt(x) for x in au.findall("AffiliationInfo/Affiliation"))
                orcid = next((_orcid(txt(x)) for x in au.findall("Identifier") if x.get("Source") == "ORCID"), "")
                detail.append({"순서": len(authors), "이름": f"{fore} {last}".strip() if last and fore else name,
                               "약칭": name, "ORCID": orcid, "소속": aff, "기관": institution(aff)})
            parts = []
            for ab in a.findall(".//Abstract/AbstractText"):
                label = ab.get("Label")
                parts.append((f"{label}: " if label else "") + txt(ab))
            pd = a.find(".//Journal/JournalIssue/PubDate")
            year = ""
            if pd is not None:
                m = re.match(r"\d{4}", pd.findtext("Year") or pd.findtext("MedlineDate") or "")
                year = m.group(0) if m else ""
            ad = art.find(".//PubmedData/History/PubMedPubDate[@PubStatus='pubmed']")
            date = "-".join(ad.findtext(k).zfill(2) for k in ("Year", "Month", "Day")) if ad is not None else ""
            doi = next((txt(x) for x in art.findall(".//PubmedData/ArticleIdList/ArticleId")
                        if x.get("IdType") == "doi"), "")
            out[pmid] = {
                "제목": txt(a.find("ArticleTitle")),
                "저자": "; ".join(authors),
                "저자수": len(authors),
                "저널": a.findtext(".//Journal/ISOAbbreviation") or a.findtext(".//Journal/Title") or "",
                "발행연도": year or date[:4],
                "발행일": date,
                "DOI": doi,
                "초록": "\n".join(parts)[:32000],        # 엑셀 셀 한도 32767자
                "_authors": detail,
            }
    return out


# ── 엑셀 읽기/쓰기 ─────────────────────────────────────────────────
def _norm(v):
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def load_db():
    """→ (rows: {pmid: {열: 값}}, queries: {데이터셋ID: {검색식, PubMed 전체건수, 조회일}})"""
    rows, queries = {}, {}
    if not DB.exists():
        return rows, queries
    wb = openpyxl.load_workbook(DB, data_only=True)
    if "논문" in wb.sheetnames:
        it = wb["논문"].iter_rows(values_only=True)
        head = [_norm(h) for h in next(it, [])]
        for r in it:
            rec = {h: _norm(v) for h, v in zip(head, r) if h}
            if rec.get("PMID"):
                rows[rec["PMID"]] = rec
    if "저자" in wb.sheetnames:
        it = wb["저자"].iter_rows(values_only=True)
        head = [_norm(h) for h in next(it, [])]
        for r in it:
            rec = {h: _norm(v) for h, v in zip(head, r) if h}
            if rec.get("PMID") in rows:
                if rec.get("소속"):
                    rec["기관"] = institution(rec["소속"])   # 사전을 고치면 재수집 없이 반영되게 읽을 때 다시 계산
                rows[rec["PMID"]].setdefault("_authors", []).append(rec)
    if "검색식" in wb.sheetnames:
        it = wb["검색식"].iter_rows(values_only=True)
        head = [_norm(h) for h in next(it, [])]
        for r in it:
            rec = {h: _norm(v) for h, v in zip(head, r) if h}
            if rec.get("데이터셋ID"):
                queries[rec["데이터셋ID"]] = rec
    return rows, queries


def load_fixes():
    """'연구자보정' 시트 → [{PMID, 이름, 연구자ID, 메모}] (사람이 적은 동명이인 보정)"""
    if not DB.exists():
        return []
    wb = openpyxl.load_workbook(DB, data_only=True)
    if "연구자보정" not in wb.sheetnames:
        return []
    it = wb["연구자보정"].iter_rows(values_only=True)
    head = [_norm(h) for h in next(it, [])]
    out = []
    for r in it:
        rec = {h: _norm(v) for h, v in zip(head, r) if h}
        if rec.get("이름") and rec.get("연구자ID"):
            out.append(rec)
    return out


def load_checks():
    """'연구자확인' 시트 → {확인ID: {열: 값}}"""
    if not DB.exists():
        return {}
    wb = openpyxl.load_workbook(DB, data_only=True)
    if "연구자확인" not in wb.sheetnames:
        return {}
    it = wb["연구자확인"].iter_rows(values_only=True)
    head = [_norm(h) for h in next(it, [])]
    out = {}
    for r in it:
        rec = {h: _norm(v) for h, v in zip(head, r) if h}
        if rec.get("확인ID"):
            out[rec["확인ID"]] = rec
    return out


def _sort_key(r):
    return (r.get("발행일") or r.get("발행연도") or "", r["PMID"])


def save_db(rows, queries, checks=None):
    fixes = load_fixes()                  # 통째로 다시 쓰므로, 사람이 적은 보정 시트는 먼저 읽어 둔다
    checks = load_checks() if checks is None else checks
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "논문"
    ws.append(COLS)
    for r in sorted(rows.values(), key=_sort_key, reverse=True):
        ws.append([int(r[c]) if c == "저자수" and str(r.get(c, "")).isdigit() else r.get(c, "") for c in COLS])
    bold, wrap = Font(bold=True, color="FFFFFF"), Alignment(vertical="top", wrap_text=True)
    for i, c in enumerate(COLS, 1):
        h = ws.cell(1, i)
        h.font, h.fill = bold, PatternFill("solid", fgColor="1F3A5F")
        ws.column_dimensions[h.column_letter].width = WIDTH.get(c, 14)
    ci = {c: i + 1 for i, c in enumerate(COLS)}
    for row in range(2, ws.max_row + 1):
        ws.row_dimensions[row].height = 45
        for col in (ci["제목"], ci["검수근거"], ci["한줄요약"], ci["초록"], ci["저자"]):
            ws.cell(row, col).alignment = wrap
        p = ws.cell(row, ci["PMID"])
        p.hyperlink, p.font = f"{PUBMED}{p.value}/", Font(color="0563C1", underline="single")
        st = ws.cell(row, ci["검수상태"])
        if st.value in FILL:
            st.fill = PatternFill("solid", fgColor=FILL[st.value])
        ap = ws.cell(row, ci["담당자승인"])
        if st.value == "확인" or ap.value:
            ap.fill = PatternFill("solid", fgColor=APPROVE_FILL.get(ap.value or "", "FFFFFF"))
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"{ws.cell(2, ci['검수상태']).column_letter}2:{ws.cell(2, ci['검수상태']).column_letter}5000")
    av = DataValidation(type="list", formula1='"' + ",".join(APPROVALS) + '"', allow_blank=True)
    ws.add_data_validation(av)
    al = ws.cell(2, ci["담당자승인"]).column_letter
    av.add(f"{al}2:{al}5000")
    ws.cell(1, ci["담당자승인"]).fill = PatternFill("solid", fgColor="B45309")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    au = wb.create_sheet("저자")          # 한 행 = 논문의 저자 1명 (fetch 가 매번 다시 만든다)
    au.append(A_COLS)
    for r in sorted(rows.values(), key=_sort_key, reverse=True):
        for a in r.get("_authors", []):
            au.append([r["PMID"]] + [a.get(c, "") for c in A_COLS[1:]])
    for i, w in enumerate([11, 6, 24, 14, 21, 70, 30], 1):
        au.cell(1, i).font, au.cell(1, i).fill = bold, PatternFill("solid", fgColor="1F3A5F")
        au.column_dimensions[au.cell(1, i).column_letter].width = w
    au.freeze_panes = "A2"
    au.auto_filter.ref = au.dimensions

    fx = wb.create_sheet("연구자보정")    # 동명이인 보정 — 사람이 적는다 (fetch 가 그대로 둔다)
    fx.append(FIX_COLS)
    for f in fixes:
        fx.append([f.get(c, "") for c in FIX_COLS])
    for i, w in enumerate([11, 24, 24, 50], 1):
        fx.cell(1, i).font, fx.cell(1, i).fill = bold, PatternFill("solid", fgColor="1F3A5F")
        fx.column_dimensions[fx.cell(1, i).column_letter].width = w
    fx.freeze_panes = "A2"

    ck = wb.create_sheet("연구자확인")    # 동명이인 확인 대상 — check 가 채우고, 사람이 '판정'을 적는다
    ck.append(CHK_COLS)
    for c in sorted(checks.values(), key=lambda c: (bool(c.get("판정")), c.get("사유", ""), c.get("이름", ""))):
        ck.append([int(c[col]) if col == "논문수" and str(c.get(col, "")).isdigit() else c.get(col, "") for col in CHK_COLS])
    for i, w in enumerate([30, 18, 20, 36, 7, 34, 26, 30, 18, 11, 11, 11, 30], 1):
        ck.cell(1, i).font, ck.cell(1, i).fill = bold, PatternFill("solid", fgColor="1F3A5F")
        ck.column_dimensions[ck.cell(1, i).column_letter].width = w
    for col in CHK_KEEP:
        ck.cell(1, CHK_COLS.index(col) + 1).fill = PatternFill("solid", fgColor="B45309")
    jc = CHK_COLS.index("판정") + 1
    for row in range(2, ck.max_row + 1):
        for col in (4, 6, 8):
            ck.cell(row, col).alignment = wrap
        if not ck.cell(row, jc).value:
            ck.cell(row, jc).fill = PatternFill("solid", fgColor="FFF4D6")
    jv = DataValidation(type="list", formula1='"' + ",".join(JUDGE) + '"', allow_blank=True)
    ck.add_data_validation(jv)
    jl = ck.cell(1, jc).column_letter
    jv.add(f"{jl}2:{jl}1000")
    ck.freeze_panes = "D2"
    ck.auto_filter.ref = ck.dimensions

    qs = wb.create_sheet("검색식")
    qs.append(Q_COLS)
    for did, q in queries.items():
        qs.append([did] + [q.get(c, "") for c in Q_COLS[1:]])
    for i, w in enumerate([14, 100, 16, 12], 1):
        qs.cell(1, i).font, qs.cell(1, i).fill = bold, PatternFill("solid", fgColor="1F3A5F")
        qs.column_dimensions[qs.cell(1, i).column_letter].width = w

    g = wb.create_sheet("안내")
    for line in [
        "논문DB — 데이터셋별 PubMed 논문 (paper_db.py 가 만든다)",
        "",
        "검수상태: 미검수 = 수집만 됨 / 확인 = 이 데이터를 실제로 썼음 / 제외 = 관련 없음 / 보류 = 초록만으로 판단 불가",
        "사이트에는 '확인'이면서 담당자승인이 '승인'인 논문만 나온다(build.py 의 SHOW_STATUS·NEED_APPROVAL). 미검수·제외·보류는 숨긴다.",
        "판정·승인은 논문 × 데이터셋 단위: 데이터셋판정(예 CD023:확인, CD006:제외)과 승인데이터셋(예 CD023)에 둘 다 있어야 그 데이터 화면에 나온다.",
        "  이미 승인된 논문에 새 데이터셋ID가 붙으면 그 데이터셋만 다시 미검수 → 검수 → 승인을 거친다.",
        "담당자승인·승인데이터셋(주황 머리글): 검수 에이전트는 건드리지 않는다. 비어 있으면 승인 대기(사이트에 안 나옴). '반려'면 어디에도 안 나옴.",
        "  한꺼번에: python paper_db.py approve <데이터셋ID>  — 그 데이터의 '확인' 논문 중 대기인 것을 모두 승인.",
        "검수상태·검수근거·검수자·검수일·메모 칸은 사람이 고쳐도 된다. 다시 수집해도 덮어쓰지 않는다.",
        "제목·저자·초록 등 나머지 칸은 다음 수집 때 PubMed 값으로 다시 채워진다(고쳐도 소용없음).",
        "행을 지우지 말 것 — 지우면 다음 수집 때 '미검수'로 다시 들어온다. 빼려면 '제외'로 바꾼다.",
        "검색식은 pubmed_queries.json 에서 고친다. '검색식' 시트는 마지막 수집 기록일 뿐이다.",
        "'저자' 시트: 논문별 저자 이름·ORCID·소속(PubMed 그대로)과 뽑아낸 기관명. 사이트의 연구자 검색에 쓰인다. 수집 때마다 다시 만들어진다.",
        "",
        "연구자 묶기(동명이인): 기본은 영문 이름이 같으면 한 사람. 같은 이름에 ORCID 가 여럿이면 기관이 겹치는 ORCID 끼리 한 사람으로 묶고",
        "  (ORCID 중복 등록이 흔하다), 기관이 안 겹치는 무리가 2개 이상일 때만 나눈다. ORCID 없는 논문은 기관이 같은 쪽에 붙이고,",
        "  못 정하면 이름 그대로 둔다. 애매한 사례는 build.py 실행 때 '[연구자 확인]'으로 출력된다.",
        "'연구자보정' 시트로 직접 고친다: 한 행 = PMID + 이름(저자 시트의 '이름' 그대로) + 연구자ID(아무 영문/한글 이름표).",
        "  연구자ID 가 같은 행끼리 한 사람이 된다. PMID 를 비우면 그 이름의 모든 논문에 적용된다.",
        "  예) 다른 사람 분리: 12345678 / Jieun Kim / jieunkim-snu   · 표기 다른 같은 사람 합치기: (빈칸) / Kim J / jieunkim",
        "  이 시트는 다시 수집해도 지워지지 않는다.",
        "",
        "'연구자확인' 시트: 동명이인일 수 있는 연구자 목록 (python paper_db.py check · fetch 때 자동으로 채움).",
        "  사유: ORCID 2개 이상 · 기관 계열 2곳 이상인데 ORCID 없음 · 표기 차이로 합쳤지만 ORCID·기관 근거 없음 · 이니셜 이름 · PubMed ORCID 오류(자동 처리, 기록만).",
        "  '판정' 칸(주황 머리글)에 같은 사람 / 다른 사람 을 고른다. 비어 있으면 '대기'.",
        "  다른 사람 → 다음 빌드부터 ORCID(없으면 기관 계열)별로 나눈다. 이니셜 이름에 같은 사람 → '대상 연구자키'로 합친다.",
        "  판정한 줄은 다시 채워도 그대로 남는다. 대상에서 빠지면 메모에 '(해당 없음)'이 붙는다.",
    ]:
        g.append([line])
    g.column_dimensions["A"].width = 120
    try:
        wb.save(DB)
    except PermissionError:
        print(f"[오류] {DB.name} 이 엑셀에서 열려 있어 저장하지 못했습니다. 닫고 다시 실행하세요.")
        sys.exit(1)


# ── 명령 ───────────────────────────────────────────────────────────
def cmd_fetch():
    conf = json.load(open(QUERIES, encoding="utf-8"))
    rows, queries = load_db()
    today = datetime.now().strftime("%Y-%m-%d")
    hits, cache = {}, {}                           # pmid → {데이터셋ID}
    for did, c in conf.items():
        if did.startswith("_") or not isinstance(c, dict) or not c.get("query", "").strip():
            continue
        q = c["query"].strip()
        if q not in cache:                          # 같은 검색식은 한 번만 조회
            cache[q] = _search(q)
            print(f"[검색] {did}: PubMed {cache[q][0]}편")
        total, ids = cache[q]
        queries[did] = {"검색식": q, "PubMed 전체건수": total, "조회일": today}
        for p in ids:
            hits.setdefault(p, set()).add(did)
        for p in c.get("exclude", []):              # 예전 방식(json exclude) 호환
            if str(p) in rows and rows[str(p)].get("검수상태") in ("", "미검수"):
                rows[str(p)].update(검수상태="제외", 검수근거="pubmed_queries.json exclude", 검수일=today)
    info = _details(list(hits)) if hits else {}
    new = 0
    for pmid, dids in hits.items():
        if pmid not in info:
            continue
        r = rows.get(pmid)
        if r is None:
            r = rows[pmid] = {"PMID": pmid, "검수상태": "미검수", "수집일": today}
            new += 1
        old = {x.strip() for x in r.get("데이터셋ID", "").split(",") if x.strip()}
        r["데이터셋ID"] = ", ".join(sorted(old | dids))
        r.update(info[pmid])                        # 서지정보만 갱신 (검수 칸은 그대로)
    save_db(rows, queries, collect_checks(rows))
    print(f"[완료] {DB.name}: 전체 {len(rows)}편 / 이번에 새로 들어온 논문 {new}편 (미검수)")
    cmd_stats(rows)


def _ids(r):
    return [x.strip() for x in r.get("데이터셋ID", "").split(",") if x.strip()]


def _verdicts(r):
    """{데이터셋ID: 확인|제외|보류} — 데이터셋판정 칸. 비어 있는 옛 행은 검수상태를 모든 데이터셋에 적용"""
    v = {}
    for part in r.get("데이터셋판정", "").split(","):
        if ":" in part:
            d, st = part.split(":", 1)
            v[d.strip()] = st.strip()
    if not v and (r.get("검수상태") or "미검수") != "미검수":
        v = {d: r["검수상태"] for d in _ids(r)}
    return v


def _set_verdicts(r, v):
    r["데이터셋판정"] = ", ".join(f"{d}:{v[d]}" for d in _ids(r) if d in v)
    sts = set(v.values())
    r["검수상태"] = "확인" if "확인" in sts else "보류" if "보류" in sts else "제외" if sts else "미검수"


def _approved(r):
    """담당자가 승인한 데이터셋 집합 — 승인데이터셋 칸. 비어 있는 옛 '승인' 행은 확인된 데이터셋 전부"""
    a = {x.strip() for x in r.get("승인데이터셋", "").split(",") if x.strip()}
    if not a and r.get("담당자승인") == "승인":
        a = {d for d, st in _verdicts(r).items() if st == "확인"}
    return a


def _unreviewed(r):
    v = _verdicts(r)
    return [d for d in _ids(r) if d not in v]


def _waiting(r):
    """확인됐지만 담당자 승인 전인 데이터셋 (반려된 논문은 제외)"""
    if r.get("담당자승인") == "반려":
        return []
    ok = _approved(r)
    return [d for d, st in _verdicts(r).items() if st == "확인" and d not in ok]


def cmd_pending(out=None):
    rows, queries = load_db()
    conf = json.load(open(QUERIES, encoding="utf-8")) if QUERIES.exists() else {}
    # 판정이 없는 데이터셋이 하나라도 있는 논문. check_datasets = 이번에 판정할 데이터셋, judged = 이미 판정된 것
    items = [{"pmid": r["PMID"], "datasets": r.get("데이터셋ID", ""), "check_datasets": ", ".join(_unreviewed(r)),
              "judged": r.get("데이터셋판정", ""), "title": r.get("제목", ""),
              "journal": r.get("저널", ""), "year": r.get("발행연도", ""), "authors": r.get("저자", ""),
              "abstract": r.get("초록", "")}
             for r in sorted(rows.values(), key=_sort_key, reverse=True) if _unreviewed(r)]
    payload = {"datasets": {d: {"query": q.get("검색식", ""), "note": conf.get(d, {}).get("note", "")}
                            for d, q in queries.items()},
               "count": len(items), "items": items}
    s = json.dumps(payload, ensure_ascii=False, indent=1)
    if out:
        Path(out).write_text(s, encoding="utf-8")
        print(f"[미검수] {len(items)}편 → {out}")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(s)


def cmd_review(path):
    decisions = json.loads(Path(path).read_text(encoding="utf-8"))
    rows, queries = load_db()
    today = datetime.now().strftime("%Y-%m-%d")
    done, bad = 0, []
    for d in decisions:
        pmid, st = str(d.get("pmid", "")).strip(), str(d.get("status", "")).strip()
        if pmid not in rows or st not in STATUSES[1:]:
            bad.append(pmid or "?")
            continue
        r = rows[pmid]
        v = _verdicts(r)
        # 판정이 적용될 데이터셋: "datasets" 를 적었으면 그것, 아니면 아직 판정 안 된 것(없으면 전부)
        targets = [x for x in (d.get("datasets") or _unreviewed(r) or _ids(r)) if x in _ids(r)]
        reason = str(d.get("reason", "")).strip()
        if v and any(x not in v for x in targets):      # 이미 검수된 논문에 새 데이터셋 → 근거를 덧붙인다
            r["검수근거"] = (r.get("검수근거", "") + f" / [{', '.join(targets)}] {reason}").strip(" /")
        else:
            r["검수근거"] = reason
        for x in targets:
            v[x] = st
        _set_verdicts(r, v)
        r.update(검수자=d.get("reviewer", "AI 검수"), 검수일=today)
        if str(d.get("summary", "")).strip() and not r.get("한줄요약"):
            r["한줄요약"] = str(d["summary"]).strip()
        done += 1
    save_db(rows, queries)
    print(f"[검수 반영] {done}편" + (f" / 건너뜀 {len(bad)}건(없는 PMID 또는 잘못된 상태): {', '.join(bad[:10])}" if bad else ""))
    cmd_stats(rows)


def cmd_nosummary(out=None):
    rows, _ = load_db()
    items = [{"pmid": r["PMID"], "title": r.get("제목", ""), "journal": r.get("저널", ""), "year": r.get("발행연도", ""),
              "abstract": r.get("초록", "")}
             for r in sorted(rows.values(), key=_sort_key, reverse=True)
             if r.get("검수상태") == "확인" and not r.get("한줄요약")]
    s = json.dumps({"count": len(items), "items": items}, ensure_ascii=False, indent=1)
    if out:
        Path(out).write_text(s, encoding="utf-8")
        print(f"[한줄요약 없음] {len(items)}편 → {out}")
    else:
        sys.stdout.reconfigure(encoding="utf-8")
        print(s)


def cmd_summary(path):
    items = json.loads(Path(path).read_text(encoding="utf-8"))
    rows, queries = load_db()
    done = 0
    for d in items:
        pmid, sm = str(d.get("pmid", "")).strip(), str(d.get("summary", "")).strip()
        if pmid in rows and sm:
            rows[pmid]["한줄요약"] = sm
            done += 1
    save_db(rows, queries)
    print(f"[한줄요약 반영] {done}편")


def collect_checks(rows=None):
    """검수 '확인' 논문 기준 확인 대상 → 기존 시트의 판정 칸을 살려 합친 {확인ID: 행}"""
    if rows is None:
        rows, _ = load_db()
    shown = [(r, []) for r in rows.values() if r.get("검수상태") == "확인"]
    old = load_checks()
    _, _, _, cases = _assign_keys(shown, load_fixes(), old)
    today = datetime.now().strftime("%Y-%m-%d")
    out = {}
    for c in cases:
        prev = old.get(c["확인ID"], {})
        row = {**c, **{k: prev[k] for k in CHK_KEEP if prev.get(k)}}
        if row.get("판정자") == "자동" and not row.get("판정일"):
            row["판정일"] = today
        out[c["확인ID"]] = row
    for cid, prev in old.items():                       # 대상에서 빠졌어도 사람이 판정한 줄은 남긴다
        if cid not in out and prev.get("판정") and prev.get("판정자") != "자동":
            memo = prev.get("메모", "")
            out[cid] = {**prev, "메모": memo if "(해당 없음)" in memo else (memo + " (해당 없음)").strip()}
    return out


def cmd_check():
    rows, queries = load_db()
    checks = collect_checks(rows)
    save_db(rows, queries, checks)
    wait = [c for c in checks.values() if not c.get("판정")]
    print(f"[연구자확인] {len(checks)}건 (판정 대기 {len(wait)}건) → {DB.name} '연구자확인' 시트")
    for c in wait:
        print(f"    · {c['이름']}: {c['사유']}")


def cmd_stats(rows=None):
    if rows is None:
        rows, _ = load_db()
    cnt = {s: 0 for s in STATUSES}
    for r in rows.values():
        cnt[r.get("검수상태") or "미검수"] = cnt.get(r.get("검수상태") or "미검수", 0) + 1
    ok = [r for r in rows.values() if r.get("검수상태") == "확인"]
    part = sum(1 for r in rows.values() if (r.get("검수상태") or "미검수") != "미검수" and _unreviewed(r))
    print("       " + " / ".join(f"{k} {v}" for k, v in cnt.items())
          + (f" (+ 새 데이터셋이 붙어 일부 미검수 {part})" if part else "")
          + f"  ·  확인 중 담당자 승인 {sum(bool(_approved(r)) for r in ok)}"
          + f" / 승인 대기 {sum(bool(_waiting(r)) for r in ok)} / 반려 {sum(r.get('담당자승인') == '반려' for r in ok)}")


def cmd_waiting():
    """담당자 승인을 기다리는 논문: 검수상태 '확인'인데 담당자승인이 빈 것"""
    rows, _ = load_db()
    wait = [r for r in sorted(rows.values(), key=_sort_key, reverse=True) if _waiting(r)]
    hold = [r for r in rows.values() if r.get("검수상태") == "보류"]
    todo = [r for r in rows.values() if (r.get("검수상태") or "미검수") == "미검수"]
    by = {}
    for r in wait:
        for d in _waiting(r):
            by[d] = by.get(d, 0) + 1
    print(f"[승인 대기] {len(wait)}편 (AI 검수 '확인', 담당자 OK 전 — 사이트에 안 나옴)"
          + (f" / 참고: 보류 {len(hold)}편 · 미검수 {len(todo)}편" if hold or todo else ""))
    for d, n in sorted(by.items(), key=lambda t: -t[1]):
        print(f"    · {d}: {n}편")
    for r in wait[:200]:
        print(f"    {r['PMID']}  [{', '.join(_waiting(r))}]  {r.get('발행연도', '')}  {r.get('제목', '')[:90]}")
        if r.get("한줄요약"):
            print(f"        └ {r['한줄요약']}")
    if wait:
        print("  OK 면: python paper_db.py approve <데이터셋ID 또는 PMID ...>   ·   빼려면: python paper_db.py reject <PMID ...>")


def cmd_approve(targets, value):
    """담당자 승인/반려 (논문 × 데이터셋 단위).
    승인: 데이터셋ID → 그 데이터셋에 '확인'인데 아직 승인 안 된 논문 전부 / PMID → 그 논문의 승인 대기 데이터셋 전부.
    반려: PMID 만 — 그 논문은 어느 데이터셋에도 안 나온다."""
    rows, queries = load_db()
    today = datetime.now().strftime("%Y-%m-%d")
    done, skip = 0, []
    for t in targets:
        if value == "반려":
            if t in rows:
                rows[t].update(담당자승인="반려", 승인일=today)
                done += 1
            else:
                skip.append(f"{t}(반려는 PMID 로만)")
            continue
        hits = [(rows[t], _waiting(rows[t]))] if t in rows else \
               [(r, [t]) for r in rows.values() if t in _waiting(r)]
        hits = [(r, ds) for r, ds in hits if ds]
        if not hits:
            skip.append(f"{t}(승인 대기 없음)")
        for r, ds in hits:
            r["승인데이터셋"] = ", ".join(x for x in _ids(r) if x in _approved(r) | set(ds))
            r.update(담당자승인="승인", 승인일=today)
            done += 1
    save_db(rows, queries)
    print(f"[담당자 {value}] {done}편" + (f" / 건너뜀: {', '.join(skip)}" if skip else ""))
    cmd_stats(rows)


# ── build.py 용 ────────────────────────────────────────────────────
def _rkey(name):
    return re.sub(r"[^a-z가-힣]", "", name.lower())


def _assign_keys(shown, fixes, checks=None):
    """저자 한 명(PMID, 순서) → 연구자키. 동명이인 처리 — '근거가 있을 때만 같은 사람':
    1) 사람 판정이 최우선: '연구자보정' 시트(PMID+이름 또는 이름만 → 연구자ID), '연구자확인' 시트의 판정
    2) 이름 표기마다 따로 시작해서, 근거가 있을 때만 합친다
         ① 같은 ORCID  ② 같은 표기(한쪽이 ORCID 없음)  ③ 같은 기관 계열(inst_family)
       서로 다른 ORCID 가 한 사람 안에 들어가는 합치기는 하지 않는다 (같은 사람이면 사람이 '같은 사람' 판정)
    3) 같은 이름키가 2명 이상으로 나뉘면 키 = 이름키-ORCID끝4자리 (ORCID 없으면 이름키-기관계열)
    4) ORCID 오류(같은 ORCID가 다른 이름에) 는 많이 붙은 쪽만 믿는다
    → (keys, notes, bad, cases)  notes = 빌드 때 출력할 문장, cases = '연구자확인' 시트에 쓸 확인 대상"""
    checks = checks or {}
    split = {c["연구자키"] for c in checks.values() if c.get("판정") == "다른 사람" and c["확인ID"].startswith("기관여러곳")}
    joins = {(c["PMID"].split(",")[0].strip(), _rkey(c["이름"])): c["대상 연구자키"] for c in checks.values()
             if c.get("판정") == "같은 사람" and c["확인ID"].startswith("이니셜") and c.get("대상 연구자키")}
    slug = lambda s: re.sub(r"[^a-z0-9가-힣-]", "", s.lower())
    by_pn = {(f["PMID"], _rkey(f["이름"])): slug(f["연구자ID"]) for f in fixes if f.get("PMID")}
    by_pn.update({k: v for k, v in joins.items() if k not in by_pn})
    by_n = {_rkey(f["이름"]): slug(f["연구자ID"]) for f in fixes if not f.get("PMID")}
    groups = {}
    for r, _ in shown:
        for a in r.get("_authors", []):
            nk = _rkey(a.get("이름", ""))
            if nk:
                groups.setdefault(nk, []).append((r["PMID"], a))
    keys, notes, cases = {}, [], []
    # ORCID 점검: 같은 ORCID가 서로 다른 이름에 붙어 있으면 PubMed 쪽 오류로 보고, 가장 많이 붙은 이름만 믿는다
    # (예: 한 논문에서 6번 저자 Hee Jin Kim의 ORCID가 7번 저자 Hee-Sun Kim에게도 붙어 있었음)
    owner = {}
    for nk, lst in groups.items():
        for _, a in lst:
            if a.get("ORCID"):
                owner.setdefault(a["ORCID"], {}).setdefault(nk, 0)
                owner[a["ORCID"]][nk] += 1
    bad = set()                                         # ORCID를 무시할 (PMID, 순서)
    for o, cnt in owner.items():
        if len(cnt) > 1:
            best = max(cnt, key=cnt.get)
            others = [g for g in cnt if g != best]
            for g in others:
                for pmid, a in groups[g]:
                    if a.get("ORCID") == o:
                        bad.add((pmid, str(a.get("순서", ""))))
            notes.append(f"ORCID {o}가 다른 이름에도 붙어 있음 → {groups[best][0][1]['이름']} 것으로 보고 "
                         + ", ".join(groups[g][0][1]["이름"] for g in others) + " 쪽은 무시 (PubMed 오류로 보임)")
            for g in others:
                pm = [pmid for pmid, a in groups[g] if a.get("ORCID") == o]
                cases.append({"확인ID": f"ORCID오류:{g}:{o}", "연구자키": g, "이름": groups[g][0][1]["이름"],
                              "사유": f"PubMed ORCID 오류 — {groups[best][0][1]['이름']}의 ORCID가 함께 붙음 (자동 무시)",
                              "ORCID": o, "PMID": ", ".join(pm), "판정": "같은 사람", "판정자": "자동",
                              "메모": "판정 불필요 — 기록용"})
    orc = lambda pmid, a: "" if (pmid, str(a.get("순서", ""))) in bad else a.get("ORCID", "")
    # 근거 기반 묶기 (이름키가 같아도 바로 합치지 않는다)
    #   단위 = (이름 표기, ORCID) — ORCID 없는 저자는 (표기, "")
    #   합치는 근거: ① 같은 ORCID  ② 같은 표기 (ORCID 없는 쪽이 있을 때)  ③ 같은 기관 계열
    #   단, 서로 다른 ORCID 가 한 사람 안에 들어가게 되는 합치기는 하지 않는다 (ORCID = 사람 고유번호)
    #   사람이 '같은 사람'으로 판정하면 그 이름키의 무리를 모두 합친다
    same_all = {c["연구자키"] for c in checks.values() if c.get("판정") == "같은 사람" and c["확인ID"].startswith("나뉨:")}
    splits = {}                                         # 이름키 → 나뉜 무리 설명 (확인 대상)
    ambig = {}                                          # 이름키 → 어느 쪽인지 애매한 ORCID 없는 표기
    for nk, lst in groups.items():
        units = {}
        for pmid, a in lst:
            uid = (a["이름"], orc(pmid, a))
            u = units.setdefault(uid, {"orc": {uid[1]} - {""}, "fam": set(), "inst": set(), "pm": set()})
            u["pm"].add(pmid)
            if a.get("기관"):
                if inst_family(a["기관"]):
                    u["fam"].add(inst_family(a["기관"]))
                u["inst"].add(a["기관"])
        par = {u: u for u in units}

        def find(u):
            while par[u] != u:
                par[u] = par[par[u]]
                u = par[u]
            return u

        def comp_orc(r):
            return set().union(*[units[u]["orc"] for u in units if find(u) == r])

        def join(x, y, force=False):
            rx, ry = find(x), find(y)
            if rx == ry:
                return True
            ox, oy = comp_orc(rx), comp_orc(ry)
            if ox and oy and not (ox & oy) and not force:
                return False                            # ORCID 가 서로 다르면 합치지 않는다
            par[ry] = rx
            return True

        ul = list(units)
        pairs = [(x, y) for i, x in enumerate(ul) for y in ul[i + 1:]]
        for x, y in pairs:                              # ① 같은 ORCID
            if units[x]["orc"] & units[y]["orc"]:
                join(x, y)
        for x, y in pairs:                              # ② 같은 표기 (한쪽이 ORCID 없음)
            if x[0] == y[0] and (not x[1] or not y[1]):
                join(x, y)
        blocked = []
        for x, y in pairs:                              # ③ 같은 기관 계열
            if units[x]["fam"] & units[y]["fam"] and not join(x, y):
                blocked.append((x, y))
        manual = any((pmid, nk) in by_pn for pmid, _ in lst) or nk in by_n
        if nk in same_all:
            for x, y in pairs:
                join(x, y, force=True)
        roots = sorted({find(u) for u in units}, key=lambda r: -sum(len(units[u]["pm"]) for u in units if find(u) == r))
        if len(roots) >= 2 and not manual and nk not in same_all:
            desc = []
            for r in roots:
                mem = [u for u in units if find(u) == r]
                o = comp_orc(r)
                desc.append(" + ".join(sorted({m[0] for m in mem})) + "("
                            + (", ".join("…" + x[-4:] for x in sorted(o)) or "ORCID 없음") + " · "
                            + (" · ".join(sorted(set().union(*[units[m]["inst"] for m in mem]))) or "기관 없음") + ")")
            why = "ORCID가 서로 다름 (기관 계열은 같음)" if blocked else "ORCID·기관 계열이 안 겹침"
            splits[nk] = [len(roots), why, " / ".join(desc), lst[0][1]["이름"]]
            notes.append(f"{lst[0][1]['이름']}: 근거가 없어 {len(roots)}명으로 나눔 — {why}")
        # ORCID 없는 표기가 서로 다른 ORCID 무리 둘 이상과 기관 계열이 겹치면 어느 쪽인지 애매하다
        for u in units:
            if not u[1] and not any(v[0] == u[0] and v[1] for v in units):   # 같은 표기의 ORCID 가 없을 때만
                hits = {find(v) for v in units if units[v]["orc"] and units[u]["fam"] & units[v]["fam"]}
                if len(hits) > 1:
                    ambig.setdefault(nk, []).append((u[0], find(u)))
        rkey = {}
        for r in roots:
            o = sorted(comp_orc(r))
            mem = [u for u in units if find(u) == r]
            fam = sorted(set().union(*[units[m]["fam"] for m in mem]))
            rkey[r] = nk if len(roots) == 1 else (f"{nk}-{o[0][-4:].lower()}" if o else
                                                  f"{nk}-{slug(fam[0])}" if fam else f"{nk}-{slug(mem[0][0])}")
        fams = {inst_family(a["기관"]) for _, a in lst if a.get("기관")} - {""}
        for pmid, a in lst:
            k = by_pn.get((pmid, nk)) or by_n.get(nk)
            if not k:
                r = find((a["이름"], orc(pmid, a)))
                k = rkey[r]
                if nk in split and not comp_orc(r) and len(fams) >= 2 and inst_family(a.get("기관")):
                    k = f"{k}-{slug(inst_family(a['기관']))}"   # ORCID 없이 '다른 사람' → 기관 계열별로
            keys[(pmid, str(a.get("순서", "")))] = k
        if nk in splits and nk in ambig:
            splits[nk][2] += " · 애매: " + ", ".join(f"{n}(ORCID 없음)은 어느 쪽인지 불확실 — 지금은 {rkey[r]} 쪽" for n, r in ambig[nk])
    # 이니셜 '같은 사람' 판정의 대상(curieahn)이 기관 계열별로 나뉘었으면(curieahn-서울대학교 …), 이니셜의 기관 계열 쪽에 붙인다
    jpos = {(pm, str(a.get("순서", ""))) for (pmid, nk) in joins for pm, a in groups.get(nk, []) if pm == pmid}
    final = {k for pos, k in keys.items() if pos not in jpos}
    for (pmid, nk), tgt in joins.items():
        if tgt in final:
            continue
        for pm, a in groups.get(nk, []):
            o = str(a.get("순서", ""))
            alt = f"{tgt}-{slug(inst_family(a.get('기관', '')))}"
            if pm == pmid and keys.get((pm, o)) == tgt and alt in final:
                keys[(pm, o)] = alt
    # 이니셜만 있는 이름(H J Kim)이 같은 약칭·같은 기관의 전체 이름(Hee Jin Kim)과 겹치면 같은 사람일 수 있다 → 확인 목록
    full = {}
    for nk, lst in groups.items():
        for pmid, a in lst:
            k = keys[(pmid, str(a.get("순서", "")))]
            if not re.match(r"^([A-Z]\.?[\s-]*)+\s+\S+$", a.get("이름", "")):
                full.setdefault(k, (a["이름"], a.get("약칭", ""), set()))[2].add(a.get("기관", ""))
    seen = set()
    for nk, lst in groups.items():
        for pmid, a in lst:
            k = keys[(pmid, str(a.get("순서", "")))]
            if k in seen or k in full or (pmid, nk) in by_pn or nk in by_n:
                continue
            if re.match(r"^([A-Z]\.?[\s-]*)+\s+\S+$", a.get("이름", "")):
                seen.add(k)
                inst = a.get("기관", "")
                cand = [v for v in full.values() if v[1] == a.get("약칭") and inst and any(i and (i in inst or inst in i) for i in v[2])]
                if len(cand) == 1:
                    notes.append(f"이니셜 이름 {a['이름']}({inst}, PMID {pmid})이 {cand[0][0]}과 같은 사람일 수 있음")
                    tk = next(kk for kk, v in full.items() if v is cand[0])
                    cases.append({"확인ID": f"이니셜:{k}:{pmid}", "연구자키": k, "이름": a["이름"],
                                  "사유": f"이니셜 이름 — {cand[0][0]}과 약칭·기관이 겹침", "기관": inst,
                                  "PMID": pmid, "대상 연구자키": tk})
    # 연구자 단위 확인 대상: ORCID 2개 이상 / 기관 계열 2곳 이상인데 ORCID 없음
    per = {}
    for nk, lst in groups.items():
        for pmid, a in lst:
            k = keys[(pmid, str(a.get("순서", "")))]
            p = per.setdefault(k, {"이름": a["이름"], "pm": [], "orc": set(), "inst": set(), "fam": set()})
            if pmid not in p["pm"]:
                p["pm"].append(pmid)
            if orc(pmid, a):
                p["orc"].add(orc(pmid, a))
            if a.get("기관"):
                p["inst"].add(a["기관"])
                if inst_family(a["기관"]):
                    p["fam"].add(inst_family(a["기관"]))
    for nk, (n, why, desc, name) in splits.items():
        ks = sorted({k for (pmid, o), k in keys.items() if any(pm == pmid for pm, _ in groups[nk]) and k.startswith(nk)})
        cases.append({"확인ID": f"나뉨:{nk}", "연구자키": nk, "이름": name, "논문수": len({pm for pm, _ in groups[nk]}),
                      "사유": f"이름이 비슷하지만 근거가 없어 {n}명으로 나눔 — {why}: {desc}",
                      "ORCID": ", ".join(sorted({orc(pm, a) for pm, a in groups[nk] if orc(pm, a)})),
                      "PMID": ", ".join(sorted({pm for pm, _ in groups[nk]})), "대상 연구자키": ", ".join(ks)})
    for k, p in per.items():
        nk = _rkey(p["이름"])
        if nk in by_n or any((pm, nk) in by_pn for pm in p["pm"]):
            continue                                    # 보정 시트로 이미 사람이 정함
        base = {"연구자키": k, "이름": p["이름"], "논문수": len(p["pm"]), "기관": " · ".join(sorted(p["inst"])),
                "ORCID": ", ".join(sorted(p["orc"])), "PMID": ", ".join(p["pm"])}
        if not p["orc"] and len(p["fam"]) >= 2:
            cases.append({**base, "확인ID": f"기관여러곳:{k}", "사유": f"ORCID 없음 · 기관 계열 {len(p['fam'])}곳 ({' / '.join(sorted(p['fam']))})"})
    return keys, notes, bad, cases


def load_for_build(dataset_ids, show_status, max_items, need_approval=True):
    """논문DB.xlsx → data.json 의 papers {데이터셋ID: {...}} 와 researchers {연구자키: {...}}
    need_approval: 담당자승인 '승인'인 논문만 (검수 에이전트의 '확인'만으로는 사이트에 안 나온다)"""
    rows, queries = load_db()
    groups, shown = {}, []
    waiting = 0
    for r in sorted(rows.values(), key=_sort_key, reverse=True):
        if (r.get("검수상태") or "미검수") not in show_status:
            continue
        waiting += bool(_waiting(r))
        if need_approval and r.get("담당자승인") == "반려":
            continue
        # 논문 × 데이터셋 단위: 그 데이터셋에 대한 판정이 show_status 이고 (승인 필요하면) 승인된 데이터셋만
        v, ok = _verdicts(r), _approved(r)
        all_ids = [x for x in _ids(r) if v.get(x) in show_status and (not need_approval or x in ok)]
        for x in all_ids:
            if x not in dataset_ids:
                print(f"[경고] 논문DB의 데이터셋ID가 목록에 없음 → 건너뜀: {x}")
        dids = [x for x in all_ids if x in dataset_ids]
        if not dids:
            continue
        shown.append((r, dids))
        for did in dids:
            groups.setdefault(did, []).append(r)

    def authors_of(r):
        return sorted(r.get("_authors", []), key=lambda a: int(a.get("순서") or 0))

    if waiting:
        print(f"[논문] 담당자 승인 대기 {waiting}편 — 사이트에 안 나옴 (python paper_db.py waiting 으로 확인)")
    checks = load_checks()
    keys, notes, bad_orcid, cases = _assign_keys(shown, load_fixes(), checks)
    pending = [c for c in cases if not (checks.get(c["확인ID"], {}).get("판정") or c.get("판정"))]
    new = [c for c in cases if c["확인ID"] not in checks]
    if pending:
        print(f"[연구자 확인] 판정 대기 {len(pending)}명 — 논문DB.xlsx '연구자확인' 시트 '판정' 칸에 같은 사람 / 다른 사람")
        for c in pending:
            print(f"    · {c['이름']}: {c['사유']}")
    if new:
        print(f"[연구자 확인] 시트에 없는 새 확인 대상 {len(new)}건 — python paper_db.py check 로 시트를 갱신하세요")
    akey = lambda r, a: keys.get((r["PMID"], str(a.get("순서", "")))) or _rkey(a.get("이름", ""))

    def item(r):
        au = authors_of(r)
        return {
            "pmid": r["PMID"],
            "title": r.get("제목", ""),
            "authors": [a.strip() for a in r.get("저자", "").split(";") if a.strip()][:6],
            "authorCount": int(r["저자수"]) if str(r.get("저자수", "")).isdigit() else 0,
            # 연구자 링크용 [정식 이름, 기관, 연구자키] — 앞 6명 + 마지막(교신) 저자
            "au": [[a["이름"], a.get("기관", ""), akey(r, a)] for a in (au[:6] + au[-1:] if len(au) > 6 else au)],
            "journal": r.get("저널", ""),
            "year": r.get("발행연도", ""),
            "doi": r.get("DOI", ""),
            "verified": r.get("검수상태") == "확인",
            "sum": r.get("한줄요약", ""),         # 한국어 1~2문장 설명 (논문 카드)
        }

    papers = {}
    for did, rs in groups.items():
        q = queries.get(did, {})
        papers[did] = {
            "query": q.get("검색식", ""),
            "total": len(rs),
            "verified": sum(1 for r in rs if r.get("검수상태") == "확인"),
            "updatedAt": q.get("조회일", ""),
            "items": [item(r) for r in rs[:max_items]],
        }

    # 연구자 색인: 사이트에 보이는 논문의 모든 저자. 묶는 기준은 _assign_keys (ORCID·연구자보정 시트)
    people = {}
    for r, dids in shown:
        au = authors_of(r)
        for i, a in enumerate(au):
            if not _rkey(a.get("이름", "")):
                continue
            k = akey(r, a)
            p = people.setdefault(k, {"name": a["이름"], "short": a.get("약칭", ""), "insts": {}, "orcids": {},
                                      "pmids": [], "datasets": set(), "first": 0, "last": 0})
            if a.get("ORCID") and (r["PMID"], str(a.get("순서", ""))) not in bad_orcid:
                p["orcids"][a["ORCID"]] = p["orcids"].get(a["ORCID"], 0) + 1
            if r["PMID"] in p["pmids"]:
                continue
            p["pmids"].append(r["PMID"])
            p["datasets"].update(dids)
            if a.get("기관"):
                p["insts"][a["기관"]] = p["insts"].get(a["기관"], 0) + 1
            p["first"] += i == 0
            p["last"] += (i == len(au) - 1 and len(au) > 1)
    researchers = {k: {"name": p["name"], "short": p["short"],
                       "insts": [x for x, _ in sorted(p["insts"].items(), key=lambda t: -t[1])][:3],
                       "pmids": p["pmids"], "datasets": sorted(p["datasets"]),
                       "first": p["first"], "last": p["last"],
                       "orcid": max(p["orcids"], key=p["orcids"].get) if p["orcids"] else ""}
                   for k, p in people.items()}
    return papers, researchers


if __name__ == "__main__":
    args = sys.argv[1:]
    if not args or args[0] not in ("fetch", "pending", "review", "stats", "nosummary", "summary", "check",
                                   "waiting", "approve", "reject"):
        print(__doc__)
        sys.exit(1)
    if args[0] == "fetch":
        cmd_fetch()
    elif args[0] == "check":
        cmd_check()
    elif args[0] == "waiting":
        cmd_waiting()
    elif args[0] in ("approve", "reject"):
        if len(args) < 2:
            print(f"사용법: python paper_db.py {args[0]} <데이터셋ID 또는 PMID> ...")
            sys.exit(1)
        cmd_approve(args[1:], "승인" if args[0] == "approve" else "반려")
    elif args[0] == "pending":
        cmd_pending(args[1] if len(args) > 1 else None)
    elif args[0] == "review":
        if len(args) < 2:
            print("사용법: python paper_db.py review 결정.json")
            sys.exit(1)
        cmd_review(args[1])
    elif args[0] == "nosummary":
        cmd_nosummary(args[1] if len(args) > 1 else None)
    elif args[0] == "summary":
        if len(args) < 2:
            print("사용법: python paper_db.py summary 요약.json")
            sys.exit(1)
        cmd_summary(args[1])
    else:
        cmd_stats()
